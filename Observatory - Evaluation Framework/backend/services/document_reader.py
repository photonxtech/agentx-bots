"""
Shared document text extraction for pdf / docx / txt.

Lives in one place because the QA generator and the RAG indexer must see the
*same* text — if they disagree, the generated ground truth can reference
content the retriever was never given, and every metric computed on top is
measuring the mismatch rather than the pipeline.
"""
import os
import logging

logger = logging.getLogger(__name__)


def _iter_docx_blocks(document):
    """Yield paragraphs and tables in true document order.

    python-docx exposes `.paragraphs` and `.tables` as separate flat lists, so
    the obvious `"\\n".join(p.text for p in doc.paragraphs)` silently discards
    every table in the file. That is not a rare edge case — report-style
    documents keep most of their substance in tables.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _read_docx(path: str) -> str:
    import docx
    from docx.table import Table

    document = docx.Document(path)
    parts: list[str] = []

    for block in _iter_docx_blocks(document):
        if isinstance(block, Table):
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells]
                # Merged cells repeat their text across the span; collapse runs
                # of identical neighbours so the chunk text is not padded out.
                deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
                line = " | ".join(c for c in deduped if c)
                if line:
                    parts.append(line)
        else:
            text = block.text.strip()
            if text:
                parts.append(text)

    return "\n".join(parts)


def _read_pdf(path: str) -> str:
    import PyPDF2
    with open(path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        return "\n".join(page.extract_text() for page in reader.pages if page.extract_text())


def read_document(path: str) -> str:
    """Extract plain text from a pdf, docx, or plain-text file."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _read_pdf(path).strip()
    if ext == ".docx":
        return _read_docx(path).strip()
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().strip()


def document_paths(session) -> list[str]:
    """Every readable document path on a session, in upload order.

    Bridges the multi-file `documents` column and the legacy single
    `document_path`, so sessions created before multi-upload still resolve.
    Missing files are skipped rather than raising — a deleted upload should not
    make the whole session unusable.
    """
    import os
    paths: list[str] = []
    for entry in (getattr(session, "documents", None) or []):
        path = (entry or {}).get("path")
        if path and path not in paths:
            paths.append(path)
    if not paths and getattr(session, "document_path", None):
        paths.append(session.document_path)
    return [p for p in paths if os.path.exists(p)]


def display_name(path: str) -> str:
    """The user's original filename, recovered from the stored path.

    Uploads are saved as `{session_uuid}_{original name}` to keep sessions from
    colliding on disk. That prefix must not leak into source attribution, where
    it would be shown to users and logged into LangSmith.
    """
    import os, re
    base = os.path.basename(path)
    return re.sub(
        r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-"
        r"[0-9a-fA-F]{4}-[0-9a-fA-F]{12}_",
        "", base,
    )


def read_documents(paths: list[str]) -> list[tuple[str, str]]:
    """Read several documents into (filename, text) pairs.

    Returned per-file rather than concatenated so downstream code can chunk each
    document separately. A chunk that straddled two unrelated files would be
    incoherent to retrieve and impossible to attribute.
    """
    out: list[tuple[str, str]] = []
    for path in paths:
        try:
            text = read_document(path)
        except Exception as e:
            logger.warning("Could not read %s: %s", path, e)
            continue
        if text:
            out.append((display_name(path), text))
    return out
