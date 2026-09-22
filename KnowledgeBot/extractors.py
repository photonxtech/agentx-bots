"""Turns a URL, a document file, or an image into plain text ready to chunk."""

import base64
import io
import os
import re
import tempfile

import easyocr
import pymupdf as fitz  # PyMuPDF's `fitz` alias is deprecated; this keeps
                         # all the fitz.* calls below unchanged either way.
                         # Used only for page rasterization and embedded-
                         # image extraction — no system binary needed
                         # (unlike poppler-based tools), which matters given
                         # past Windows pain with system-level OCR deps.
import pandas as pd
import pdfplumber
import requests
import trafilatura
from docx import Document
from lxml import html as lxml_html
from pptx import Presentation

# Loaded once and reused across calls — EasyOCR's model load is the slow
# part, so we don't want to repeat it per image. English by default; add
# more language codes here if needed, e.g. easyocr.Reader(["en", "es"]).
_ocr_reader = None


def _get_ocr_reader():
    global _ocr_reader
    if _ocr_reader is None:
        _ocr_reader = easyocr.Reader(["en"], gpu=False)
    return _ocr_reader


# A real browser User-Agent — trafilatura's default fetcher identifies
# itself in a way that some sites' bot-protection (Cloudflare, WAFs, etc.)
# silently blocks or redirects, even when the page needs no JavaScript at
# all and would otherwise scrape fine. Fetching manually with a normal
# browser UA avoids a lot of false "couldn't read this" failures.
_BROWSER_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}


def extract_from_url(
    url: str, groq_client=None, vision_model: str = None, max_images: int = 4
) -> tuple[str, str]:
    """Returns (title, text) for a web page.

    Text always includes the clean article content. If groq_client and
    vision_model are given, up to `max_images` meaningful images on the
    page (icons/logos/etc. filtered out) are also OCR'd and captioned,
    with the descriptions appended to the returned text — so a question
    about the page can draw on what's shown, not just what's written.
    """
    try:
        resp = requests.get(url, headers=_BROWSER_HEADERS, timeout=15)
        resp.raise_for_status()
        downloaded = resp.text
    except requests.exceptions.RequestException as direct_error:
        # Direct fetch got blocked (403, bot-protection, etc.) — no header
        # tweak reliably gets past real bot-protection services, so fall
        # back to a reader-proxy that fetches the page server-side and
        # hands back clean text. This won't extract images (we don't get
        # the raw HTML from it), but it's better than failing outright.
        return _extract_via_reader_proxy(url, direct_error)

    text = trafilatura.extract(downloaded, include_comments=False, include_tables=True)
    metadata = trafilatura.extract_metadata(downloaded)
    title = (metadata.title if metadata and metadata.title else url)

    if not text:
        raise ValueError(f"Could not extract readable content from {url}")

    if groq_client and vision_model:
        image_urls = _find_content_images(downloaded, url, max_images)
        descriptions = []
        for img_url in image_urls:
            try:
                desc = _describe_remote_image(img_url, groq_client, vision_model)
                if desc:
                    descriptions.append(desc)
            except Exception:
                # A single bad/unreachable image shouldn't fail the whole
                # page ingestion — just skip it and move on.
                continue

        if descriptions:
            text += "\n\nImages found on this page:\n" + "\n\n".join(descriptions)

    return title, text


def _extract_via_reader_proxy(url: str, direct_error: Exception) -> tuple[str, str]:
    """Fallback for pages that block direct scraping. r.jina.ai fetches the
    page server-side (from its own IPs, real browser rendering) and returns
    clean text — this gets past bot-protection that blocks us directly,
    at the cost of not being able to pull images out of the page.
    """
    try:
        resp = requests.get(f"https://r.jina.ai/{url}", timeout=25)
        resp.raise_for_status()
    except requests.exceptions.RequestException as proxy_error:
        raise ValueError(
            f"Could not fetch {url} (direct: {direct_error}; reader-proxy fallback also failed: {proxy_error})"
        )

    text = resp.text.strip()
    if not text:
        raise ValueError(f"Could not extract any content from {url} via direct fetch or fallback")

    # r.jina.ai typically starts its output with a "Title: ..." line.
    title = url
    first_line = text.splitlines()[0] if text else ""
    if first_line.lower().startswith("title:"):
        title = first_line.split(":", 1)[1].strip()

    return title, text


_JUNK_IMAGE_PATTERN = re.compile(
    r"(icon|logo|avatar|sprite|tracking|badge|button|spacer|pixel\.)", re.IGNORECASE
)
_MIN_IMAGE_DIMENSION = 150  # px, only enforced when the <img> tag declares it


def _find_content_images(raw_html: str, base_url: str, max_images: int) -> list[str]:
    """Pulls likely-meaningful image URLs out of a page's HTML, filtering
    out obvious junk (icons/logos/tracking pixels) so we don't waste OCR
    and vision-model calls on decoration.
    """
    try:
        tree = lxml_html.fromstring(raw_html)
    except Exception:
        return []

    candidates = []
    for img in tree.xpath("//img"):
        src = img.get("src") or img.get("data-src")
        if not src:
            continue

        full_url = requests.compat.urljoin(base_url, src)

        if _JUNK_IMAGE_PATTERN.search(full_url):
            continue

        width, height = img.get("width"), img.get("height")
        try:
            if width and int(width) < _MIN_IMAGE_DIMENSION:
                continue
            if height and int(height) < _MIN_IMAGE_DIMENSION:
                continue
        except ValueError:
            pass  # non-numeric width/height (e.g. "100%"), can't filter on it, allow through

        candidates.append(full_url)
        if len(candidates) >= max_images:
            break

    return candidates


def _describe_remote_image(image_url: str, groq_client, vision_model: str) -> str:
    """Downloads one image from a page and runs it through the same
    OCR + captioning pipeline used for uploaded images.
    """
    resp = requests.get(image_url, timeout=15)
    resp.raise_for_status()

    content_type = resp.headers.get("Content-Type", "")
    if "image" not in content_type:
        return ""
    if len(resp.content) > 8 * 1024 * 1024:  # skip anything unreasonably large
        return ""

    ext = content_type.split("/")[-1].split(";")[0].lower() or "png"
    if ext not in ("png", "jpg", "jpeg", "gif", "webp"):
        ext = "png"

    with tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False) as tmp:
        tmp.write(resp.content)
        tmp_path = tmp.name

    try:
        ocr_text = ocr_image(tmp_path)
        caption = caption_image(tmp_path, groq_client, vision_model)
        combined = f"- {caption}"
        if ocr_text:
            combined += f" (text in image: {ocr_text})"
        return combined
    finally:
        os.remove(tmp_path)


def extract_from_pdf(file_path: str, groq_client=None, vision_model: str = None) -> str:
    """Routes based on what kind of PDF this actually is, instead of
    always running the full (slow) image pipeline regardless of need:

    - Text-based PDF, no embedded images: plain text extraction only.
    - Text-based PDF WITH embedded images (charts/diagrams): text
      extraction, plus OCR + vision captioning on just those images.
    - Scanned PDF (no real text layer at all): renders each page as an
      image and OCRs it, since there's no text to extract directly.
    """
    if _is_scanned_pdf(file_path):
        return _extract_scanned_pdf_via_ocr(file_path)

    parts = []
    has_embedded_images = False
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
            if page.images:
                has_embedded_images = True
    text = "\n\n".join(parts)

    if has_embedded_images and groq_client and vision_model:
        descriptions = _extract_and_caption_pdf_images(file_path, groq_client, vision_model)
        if descriptions:
            text += "\n\nImages/charts found in this PDF:\n" + "\n\n".join(descriptions)

    return text


def _is_scanned_pdf(file_path: str, sample_pages: int = 3) -> bool:
    """Checks a handful of pages for a real, extractable text layer. If
    none of the sampled pages have meaningful text, this is almost
    certainly a scanned document (pages are just images of pages) rather
    than a text-based PDF, so it needs OCR instead of direct extraction.
    """
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages[:sample_pages]:
            text = (page.extract_text() or "").strip()
            if len(text) > 20:  # a real text layer, not stray noise
                return False
    return True


def _render_pdf_pages_to_images(file_path: str, dpi: int = 200) -> list[str]:
    """Rasterizes each page to a temp PNG for OCR, via PyMuPDF — no
    poppler/Ghostscript install needed, unlike pdf2image.
    """
    doc = fitz.open(file_path)
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    paths = []
    try:
        for page in doc:
            pix = page.get_pixmap(matrix=matrix)
            tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
            pix.save(tmp.name)
            paths.append(tmp.name)
    finally:
        doc.close()
    return paths


def _extract_scanned_pdf_via_ocr(file_path: str) -> str:
    image_paths = _render_pdf_pages_to_images(file_path)
    try:
        parts = []
        for i, img_path in enumerate(image_paths, start=1):
            text = ocr_image(img_path)
            if text:
                parts.append(f"Page {i}:\n{text}")
        return "\n\n".join(parts)
    finally:
        for p in image_paths:
            try:
                os.remove(p)
            except OSError:
                pass


def _extract_and_caption_pdf_images(
    file_path: str, groq_client, vision_model: str, max_images: int = 4
) -> list[str]:
    """Pulls embedded images (charts/diagrams) out of a text-based PDF and
    runs them through the same OCR + captioning pipeline used elsewhere,
    so a question about a chart's content is actually answerable.
    """
    doc = fitz.open(file_path)
    descriptions = []
    try:
        for page_index in range(len(doc)):
            if len(descriptions) >= max_images:
                break
            for img in doc.get_page_images(page_index):
                if len(descriptions) >= max_images:
                    break
                xref = img[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]

                if len(image_bytes) < 3000:
                    # Crude size-based junk filter — very small embedded
                    # images are almost always icons/bullets/decoration,
                    # not real charts, not worth a captioning call.
                    continue

                ext = base_image.get("ext", "png")
                with tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False) as tmp:
                    tmp.write(image_bytes)
                    tmp_path = tmp.name

                try:
                    ocr_text = ocr_image(tmp_path)
                    caption = caption_image(tmp_path, groq_client, vision_model)
                    combined = f"- {caption}"
                    if ocr_text:
                        combined += f" (text in image: {ocr_text})"
                    descriptions.append(combined)
                except Exception:
                    continue  # a bad embedded image shouldn't fail the whole PDF
                finally:
                    os.remove(tmp_path)
    finally:
        doc.close()
    return descriptions


def extract_from_docx(file_path: str) -> str:
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # doc.paragraphs only walks top-level text — table cells need separate
    # handling or they're silently dropped.
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    return "\n\n".join(parts)


def extract_from_csv(file_path: str) -> str:
    df = pd.read_csv(file_path)
    header = f"Columns: {', '.join(str(c) for c in df.columns)}\nRows: {len(df)}\n\n"
    return header + df.to_csv(index=False)


def extract_from_excel(file_path: str) -> str:
    """Handles multi-sheet workbooks — each sheet extracted separately and
    labeled, since a question might be about any one of them.
    """
    xls = pd.ExcelFile(file_path)
    parts = []
    for sheet_name in xls.sheet_names:
        df = xls.parse(sheet_name)
        header = (
            f"Sheet: {sheet_name}\n"
            f"Columns: {', '.join(str(c) for c in df.columns)}\n"
            f"Rows: {len(df)}\n\n"
        )
        parts.append(header + df.to_csv(index=False))
    return "\n\n---\n\n".join(parts)


def extract_from_pptx(file_path: str) -> str:
    """Pulls text from slide bodies, tables, and speaker notes."""
    prs = Presentation(file_path)
    parts = []

    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip(" |"):
                        lines.append(row_text)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")

        if lines:
            parts.append(f"Slide {i}:\n" + "\n".join(lines))

    return "\n\n".join(parts)


def extract_from_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def ocr_image(file_path: str) -> str:
    """Pulls any literal text out of an image."""
    reader = _get_ocr_reader()
    results = reader.readtext(file_path, detail=0)  # detail=0 -> just the text strings
    return " ".join(results).strip()


def caption_image(file_path: str, groq_client, vision_model: str) -> str:
    """Uses a Groq vision model to describe what's actually in the image."""
    with open(file_path, "rb") as f:
        b64 = base64.b64encode(f.read()).decode("utf-8")

    ext = os.path.splitext(file_path)[1].lstrip(".").lower() or "png"
    mime = "jpeg" if ext in ("jpg", "jpeg") else ext

    response = groq_client.chat.completions.create(
        model=vision_model,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "Describe this image in detail: what it shows, any "
                            "charts/diagrams/UI elements, and the key information "
                            "someone would want to search for later."
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/{mime};base64,{b64}"},
                    },
                ],
            }
        ],
        max_tokens=500,
    )
    return response.choices[0].message.content.strip()


def extract_from_image(file_path: str, groq_client, vision_model: str) -> str:
    """Combines OCR text with a vision-model caption into one text blob."""
    ocr_text = ocr_image(file_path)
    caption = caption_image(file_path, groq_client, vision_model)

    combined = f"Image description: {caption}"
    if ocr_text:
        combined += f"\n\nText found in image (OCR): {ocr_text}"
    return combined


def download_slack_file(file_url: str, bot_token: str, dest_path: str) -> str:
    headers = {"Authorization": f"Bearer {bot_token}"}
    resp = requests.get(file_url, headers=headers, timeout=30)
    resp.raise_for_status()
    with open(dest_path, "wb") as f:
        f.write(resp.content)
    return dest_path